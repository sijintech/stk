"""Real XLSX files and typed JSON preserve scientific values without execution."""
from datetime import date
import hashlib
import io
import json
import os
from pathlib import Path
from uuid import uuid4
import zipfile

import openpyxl
import pytest

from suan.project import ProjectStore, ProjectError, RevisionConflict
from suan.project.table_io import TableIO


def uid():
    return str(uuid4())


@pytest.fixture
def store(tmp_path):
    return ProjectStore.create(tmp_path / "project", "Table exchange")


@pytest.fixture
def typed(store):
    table_id, first, second, third, choice = [uid() for _ in range(5)]
    choice_options = {"version": 1, "choices": [{"id": choice, "label": "Ready", "color": "#123456"}]}
    kinds = ["text", "integer", "number", "boolean", "json", "date", "url", "single_select", "multi_select", "object"]
    fields = [{"id": uid(), "name": k, "type": k,
               "unit": "K" if k == "number" else None,
               "options": choice_options if k.endswith("select") else {"version": 1, "kinds": ["record"]}
               if k == "object" else {"version": 1}} for k in kinds]
    values = ['=HYPERLINK("https://example.org","literal")\n中文', 9223372036854775807, 300.125, True,
              {"嵌套": [None, {"x": "multi\nline"}], "int64": 9007199254740993}, "2024-02-29",
              "https://example.org/a?q=1", choice, [choice],
              {"version": 1, "project_id": store.info()["id"], "kind": "record", "id": second, "table_id": table_id}]
    commands = [{"op": "create_table", "id": table_id, "name": "全类型 Cases"}]
    commands += [{"op": "add_field", "table_id": table_id, **field} for field in fields]
    commands += [{"op": "add_record", "table_id": table_id, "id": identity} for identity in (first, second, third)]
    commands += [{"op": "set_cell", "table_id": table_id, "record_id": first, "field_id": f["id"], "value": value}
                 for f, value in zip(fields, values)]
    commands += [{"op": "set_cell", "table_id": table_id, "record_id": second, "field_id": f["id"], "value": None}
                 for f in fields]
    store.apply(commands, expected_revision=0)
    return table_id, fields, values, (first, second, third)


def matrix(store, table_id):
    table = next(t for t in store.snapshot()["tables"] if t["id"] == table_id)
    return [[r["values"].get(f["id"], "<unset>") for f in table["fields"]] for r in table["records"]]


def workbook(path, *, formula=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "First"
    ws.append(["Name", "Count"]); ws.append(["001", 7])
    other = wb.create_sheet("实验数据")
    other.append(["Text", "Count", "Temperature", "Enabled", "Date", "Details"])
    other.append(["=1+2" if formula else "中文", "9223372036854775807", 301.25, True,
                  date(2024, 2, 29), '{"x":[1,null]}'])
    wb.save(path); wb.close()
    return path


def test_json_all_types_null_unset_schema_new_ids_and_one_undo(store, typed, tmp_path):
    table_id, fields, values, identities = typed
    service = TableIO(store)
    path = tmp_path / "table.json"
    before = store.snapshot()
    exported = service.export_file(table_id, path, format="json", expected_revision=1)
    assert store.snapshot() == before and exported["sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()
    raw = json.loads(path.read_text())
    assert raw["format"] == "stk.table" and raw["version"] == 1 and raw["mode"] == "evaluated_values"
    assert raw["records"][2]["values"] == {} and set(raw["records"][1]["values"].values()) == {None}
    preview = service.preview(path, format="json", expected_revision=1, limit=1)
    assert preview["schema_present"] and preview["total_rows"] == 3 and preview["truncated"]
    assert preview["rows"] == [values] and preview["columns"][7]["options"] == fields[7]["options"]
    result = service.import_file(path, format="json", name="Copy", expected_revision=1,
                                 expected_sha256=preview["source_sha256"])
    assert result["revision"] == 2 and result["table_id"] != table_id
    assert not set(result["field_ids"]) & {f["id"] for f in fields}
    assert not set(result["record_ids"]) & set(identities)
    copied = store.snapshot()["tables"][-1]
    assert [{k: f[k] for k in ("name", "type", "unit", "options")} for f in copied["fields"]] == [
        {k: f[k] for k in ("name", "type", "unit", "options")} for f in fields]
    assert matrix(store, result["table_id"]) == [values, [None] * 10, ["<unset>"] * 10]
    store.undo(expected_revision=2)
    assert store.snapshot()["tables"] == before["tables"]
    store.redo(expected_revision=3)
    assert matrix(store, result["table_id"])[0] == values


def test_xlsx_export_is_real_inert_formula_text_and_exact_int64(store, typed, tmp_path):
    table_id, fields, values, _ = typed
    service = TableIO(store)
    path = tmp_path / "table.xlsx"
    result = service.export_file(table_id, path, format="xlsx", expected_revision=1)
    assert path.read_bytes().startswith(b"PK") and result["rows"] == 3
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb.active
    assert ws["A2"].value == values[0] and ws["A2"].data_type == "s"
    assert ws["B2"].value == "9223372036854775807" and ws["B2"].data_type == "s"
    assert ws["D2"].value is True and ws["C2"].value == 300.125
    assert json.loads(ws["E2"].value) == values[4]
    assert not wb._external_links and not any(c.data_type == "f" for row in ws for c in row)
    wb.close()
    preview = service.preview(path, format="xlsx", expected_revision=1)
    assert preview["format"] == "xlsx" and preview["sheet"] == "Table"
    assert preview["sheets"] == [{"name": "Table"}] and not preview["schema_present"]
    assert preview["rows"][0][1] == "9223372036854775807"
    assert preview["rows"][0][0] == values[0]


def test_xlsx_sheet_preview_typed_import_and_entire_batch_undo(store, tmp_path):
    path = workbook(tmp_path / "input.xlsx")
    service = TableIO(store)
    first = service.preview(path, format="xlsx", expected_revision=0)
    assert first["sheet"] == "First" and first["rows"] == [["001", "7"]]
    preview = service.preview(path, format="xlsx", sheet="实验数据", expected_revision=0)
    assert preview["sheets"] == [{"name": "First"}, {"name": "实验数据"}]
    assert preview["rows"] == [["中文", "9223372036854775807", "301.25", "true", "2024-02-29", '{"x":[1,null]}']]
    assert store.info()["revision"] == 0 and store.history() == []
    result = service.import_file(path, format="xlsx", name="Imported", sheet="实验数据", expected_revision=0,
        types={"Count": "integer", "Temperature": "number", "Enabled": "boolean", "Date": "date", "Details": "json"},
        units={"Temperature": "K"}, expected_sha256=preview["source_sha256"])
    assert matrix(store, result["table_id"]) == [["中文", 9223372036854775807, 301.25, True, "2024-02-29", {"x": [1, None]}]]
    assert store.snapshot()["tables"][0]["fields"][2]["unit"] == "K"
    store.undo(expected_revision=1)
    assert store.snapshot()["tables"] == []


def test_csv_wrapper_preserves_existing_contract(store, tmp_path):
    source = tmp_path / "in.csv"; source.write_text("A,B\n001,7\n", encoding="utf-8")
    service = TableIO(store)
    preview = service.preview(source, format="csv", expected_revision=0)
    assert preview["rows"] == [["001", "7"]] and preview["format"] == "csv"
    result = service.import_file(source, format="csv", expected_revision=0, name="CSV", types={"B": "integer"},
                                 expected_sha256=preview["source_sha256"])
    assert matrix(store, result["table_id"]) == [["001", 7]]
    exported = service.export_file(result["table_id"], tmp_path / "out.tsv", format="csv", delimiter="\t", expected_revision=1)
    assert Path(exported["path"]).read_text(encoding="utf-8-sig") == "A\tB\n001\t7\n"


@pytest.mark.parametrize("format", ["xlsx", "json"])
def test_sha_change_and_stale_revision_reject_without_partial_import(store, typed, tmp_path, format):
    table_id, _, _, _ = typed
    service = TableIO(store); path = tmp_path / f"source.{format}"
    service.export_file(table_id, path, format=format, expected_revision=1)
    preview = service.preview(path, format=format, expected_revision=1)
    if format == "json":
        path.write_text(path.read_text() + " ")
    else:
        wb = openpyxl.load_workbook(path); wb.active["A2"] = "changed"; wb.save(path); wb.close()
    before = store.snapshot()
    with pytest.raises(RevisionConflict, match="source changed"):
        service.import_file(path, format=format, name="Stale bytes", expected_revision=1,
                            expected_sha256=preview["source_sha256"])
    with pytest.raises(RevisionConflict):
        service.preview(path, format=format, expected_revision=0)
    with pytest.raises(RevisionConflict):
        service.export_file(table_id, tmp_path / f"stale.{format}", format=format, expected_revision=0)
    assert store.snapshot() == before


@pytest.mark.parametrize("format", ["xlsx", "json"])
def test_final_apply_cas_rejects_concurrent_writer(store, tmp_path, format, monkeypatch):
    service = TableIO(store)
    source = tmp_path / "a.csv"; source.write_text("A\nx\n")
    created = service.import_file(source, format="csv", name="Original", expected_revision=0)
    path = tmp_path / f"source.{format}"
    service.export_file(created["table_id"], path, format=format, expected_revision=1)
    original = store.apply
    def concurrent(commands, *, expected_revision):
        original([{"op": "create_table", "name": "Concurrent writer"}], expected_revision=1)
        return original(commands, expected_revision=expected_revision)
    monkeypatch.setattr(store, "apply", concurrent)
    with pytest.raises(RevisionConflict):
        service.import_file(path, format=format, name="Stale import", expected_revision=1)
    assert [t["name"] for t in store.snapshot()["tables"]] == ["Original", "Concurrent writer"]


@pytest.mark.parametrize("format", ["xlsx", "json"])
def test_export_create_only_race_leaves_destination_and_no_temporary_file(store, typed, tmp_path, format, monkeypatch):
    table_id, _, _, _ = typed
    destination = tmp_path / f"out.{format}"
    link = os.link
    def raced(source, target):
        Path(target).write_bytes(b"other owner's file")
        return link(source, target)
    monkeypatch.setattr(os, "link", raced)
    with pytest.raises(FileExistsError):
        TableIO(store).export_file(table_id, destination, format=format, expected_revision=1)
    assert destination.read_bytes() == b"other owner's file"
    assert not list(tmp_path.glob(".stk-table-*"))


def test_json_cross_project_object_refs_reject_but_value_only_tables_are_portable(store, typed, tmp_path):
    table_id, _, _, _ = typed
    path = tmp_path / "data.json"
    TableIO(store).export_file(table_id, path, format="json", expected_revision=1)
    other = ProjectStore.create(tmp_path / "other", "Other")
    with pytest.raises(ProjectError, match="project"):
        TableIO(other).import_file(path, format="json", name="Wrong source", expected_revision=0)
    assert other.history() == []
    doc = json.loads(path.read_text()); object_id = doc["fields"][-1]["id"]
    for row in doc["records"]:
        row["values"].pop(object_id, None)
    path.write_text(json.dumps(doc))
    result = TableIO(other).import_file(path, format="json", name="Values", expected_revision=0)
    assert result["rows"] == 3 and matrix(other, result["table_id"])[0][1] == 9223372036854775807


def test_filtered_sorted_visible_columns_export_identically_to_json_and_excel(store, typed, tmp_path):
    table_id, fields, values, identities = typed
    definition = {"version": 1, "kind": "grid", "columns": [
        {"field_id": fields[2]["id"]}, {"field_id": fields[0]["id"]}, {"field_id": fields[1]["id"], "hidden": True}],
        "filters": {"field_id": fields[2]["id"], "op": "gt", "value": 300},
        "sorts": [{"field_id": fields[2]["id"], "direction": "desc"}]}
    for format in ("json", "xlsx"):
        path = tmp_path / f"filtered.{format}"
        result = TableIO(store).export_file(table_id, path, format=format, expected_revision=1, definition=definition)
        assert result["rows"] == 1 and result["columns"] == 2
        preview = TableIO(store).preview(path, format=format, expected_revision=1)
        assert [f["name"] for f in preview["columns"]] == ["number", "text"]
        assert preview["rows"] == [[300.125 if format == "json" else "300.125", values[0]]]
    assert store.info()["revision"] == 1


def test_json_materializes_successful_formulas_and_rejects_errors(store, tmp_path):
    source = tmp_path / "a.csv"; source.write_text("A,B\n3,0\n")
    service = TableIO(store)
    r = service.import_file(source, format="csv", name="Formulas", types={"A": "integer", "B": "integer"}, expected_revision=0)
    table_id, record_id = r["table_id"], r["record_ids"][0]
    a, b = r["field_ids"]
    command = {"op": "set_expression", "table_id": table_id, "record_id": record_id, "field_id": b,
               "expression": "x * 2", "bindings": {"x": {"record_id": record_id, "field_id": a}}}
    store.apply([command], expected_revision=1)
    path = tmp_path / "formula.json"
    service.export_file(table_id, path, format="json", expected_revision=2)
    document = json.loads(path.read_text())
    assert document["records"][0]["values"][b] == 6 and "definitions" not in document["records"][0]
    copy = service.import_file(path, format="json", name="Materialized", expected_revision=2)
    assert "definitions" not in store.snapshot()["tables"][-1]["records"][0]
    command["expression"] = "x / 0"
    store.apply([command], expected_revision=3)
    for format in ("json", "xlsx"):
        destination = tmp_path / f"error.{format}"
        with pytest.raises(ProjectError, match="formula error"):
            service.export_file(table_id, destination, format=format, expected_revision=4)
        assert not destination.exists()


def repack(path, replacements):
    output = io.BytesIO()
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for member in source.infolist():
            if member.filename not in replacements:
                target.writestr(member.filename, source.read(member))
        for name, data in replacements.items():
            target.writestr(name, data)
    path.write_bytes(output.getvalue())


@pytest.mark.parametrize("attack", ["formula", "vba", "external", "entity", "expanded", "dimensions", "duplicate"])
def test_excel_unsafe_or_unbounded_inputs_reject_without_edit(store, tmp_path, attack):
    path = workbook(tmp_path / "unsafe.xlsx", formula=attack == "formula")
    if attack == "vba":
        repack(path, {"xl/vbaProject.bin": b"not executed"})
    elif attack == "external":
        repack(path, {"xl/_rels/evil.rels": '<Relationships><Relationship TargetMode="External" Target="https://example.org"/></Relationships>'})
    elif attack == "entity":
        repack(path, {"evil.xml": '<!DOCTYPE root [<!ENTITY x "expanded">]><root>&x;</root>'})
    elif attack == "expanded":
        repack(path, {"oversized.xml": b"x" * (32 * 1024 * 1024 + 1)})
    elif attack == "dimensions":
        wb = openpyxl.load_workbook(path); wb["实验数据"]["BM10002"] = "far"; wb.save(path); wb.close()
    elif attack == "duplicate":
        with zipfile.ZipFile(path, "a") as archive, pytest.warns(UserWarning, match="Duplicate"):
            archive.writestr("[Content_Types].xml", "<duplicate/>")
    with pytest.raises(ProjectError):
        TableIO(store).preview(path, format="xlsx", sheet="实验数据", expected_revision=0)
    with pytest.raises(ProjectError):
        TableIO(store).import_file(path, format="xlsx", sheet="实验数据", name="Unsafe", expected_revision=0)
    assert store.info()["revision"] == 0 and store.history() == []


@pytest.mark.parametrize("change", ["version", "duplicate_fields", "unknown_values", "duplicate_records", "invalid_date", "overrides", "nonfinite", "duplicate_keys", "deep"])
def test_invalid_json_rejects_atomically(store, typed, tmp_path, change):
    table_id, fields, _, _ = typed
    path = tmp_path / "invalid.json"
    service = TableIO(store)
    service.export_file(table_id, path, format="json", expected_revision=1)
    doc = json.loads(path.read_text()); kwargs = {}
    if change == "version": doc["version"] = True
    elif change == "duplicate_fields": doc["fields"][1]["id"] = doc["fields"][0]["id"]
    elif change == "unknown_values": doc["records"][0]["values"][uid()] = "extra"
    elif change == "duplicate_records": doc["records"][1]["id"] = doc["records"][0]["id"]
    elif change == "invalid_date": doc["records"][0]["values"][fields[5]["id"]] = "2023-02-29"
    elif change == "overrides": kwargs["types"] = {"text": "integer"}
    path.write_text(json.dumps(doc), encoding="utf-8")
    if change == "nonfinite": path.write_text(path.read_text().replace("300.125", "NaN"))
    elif change == "duplicate_keys": path.write_text(path.read_text().replace('"version": 1', '"version":1,"version":1', 1))
    elif change == "deep": path.write_text("[" * 65 + "0" + "]" * 65)
    before = store.snapshot()
    with pytest.raises(ProjectError):
        service.import_file(path, format="json", name="Bad", expected_revision=1, **kwargs)
    assert store.snapshot() == before and store.info()["revision"] == 1


def test_xlsx_atomic_batch_limit_preview_and_nonregular_sources(store, tmp_path):
    path = tmp_path / "many.xlsx"
    wb = openpyxl.Workbook(); ws = wb.active; ws.append(["A", "B"])
    for i in range(333): ws.append([i, i])
    wb.save(path); wb.close()
    service = TableIO(store)
    preview = service.preview(path, format="xlsx", expected_revision=0)
    assert preview["command_count"] == 1002 and not preview["can_import"]
    with pytest.raises(ProjectError, match="1000"):
        service.import_file(path, format="xlsx", name="Too many", expected_revision=0)
    if os.name == "posix":
        fifo = tmp_path / "fifo"; os.mkfifo(fifo)
        with pytest.raises(ProjectError, match="regular"):
            service.preview(fifo, format="json", expected_revision=0)
    assert store.info()["revision"] == 0


@pytest.mark.parametrize("format", ["xlsx", "json"])
def test_export_rechecks_revision_before_publication(store, typed, tmp_path, format, monkeypatch):
    table_id, _, _, _ = typed
    service = TableIO(store)
    original = service._table
    def racing(*args):
        value = original(*args)
        store.apply([{"op": "create_table", "name": "Concurrent"}], expected_revision=1)
        return value
    monkeypatch.setattr(service, "_table", racing)
    path = tmp_path / f"out.{format}"
    with pytest.raises(RevisionConflict):
        service.export_file(table_id, path, format=format, expected_revision=1)
    assert not path.exists()


@pytest.mark.parametrize("value", ["x" * 32768, "contains\x00control", "contains\x0bvertical-tab"])
def test_excel_rejects_unrepresentable_text_instead_of_silent_truncation(store, tmp_path, value):
    ids = [uid() for _ in range(3)]
    store.apply([{"op": "create_table", "id": ids[0], "name": "Text"},
                 {"op": "add_field", "table_id": ids[0], "id": ids[1], "name": "A", "type": "text"},
                 {"op": "add_record", "table_id": ids[0], "id": ids[2]},
                 {"op": "set_cell", "table_id": ids[0], "record_id": ids[2], "field_id": ids[1], "value": value}], expected_revision=0)
    destination = tmp_path / "cannot.xlsx"
    with pytest.raises(ProjectError, match="32767"):
        TableIO(store).export_file(ids[0], destination, format="xlsx", expected_revision=1)
    assert not destination.exists() and store.info()["revision"] == 1


@pytest.mark.parametrize("format", ["json", "xlsx"])
def test_command_byte_budget_preview_matches_atomic_import_rejection(store, tmp_path, format):
    source = tmp_path / "large.csv"
    source.write_text("A\n" + ("x" * 30000 + "\n") * 40)
    # Existing CSV compatibility permits this 82-command table; new rich exchange
    # must report and enforce its own 1 MiB materialized-command budget.
    r = store.csv.import_file(source, name="Source", expected_revision=0)
    service = TableIO(store)
    path = tmp_path / f"large.{format}"
    service.export_file(r["table_id"], path, format=format, expected_revision=1)
    preview = service.preview(path, format=format, expected_revision=1)
    assert preview["command_count"] == 82 and not preview["can_import"]
    with pytest.raises(ProjectError, match="1 MiB"):
        service.import_file(path, format=format, name="Too large", expected_revision=1)
    assert len(store.snapshot()["tables"]) == 1 and store.info()["revision"] == 1


def test_core_json_import_into_format13_requires_no_implicit_upgrade(store, tmp_path):
    import sqlite3
    source = tmp_path / "core.csv"; source.write_text("A\n7\n")
    r = store.csv.import_file(source, name="Core", types={"A": "integer"}, expected_revision=0)
    json_path = tmp_path / "core.json"
    TableIO(store).export_file(r["table_id"], json_path, format="json", expected_revision=1)
    other = ProjectStore.create(tmp_path / "legacy", "Legacy")
    with sqlite3.connect(other.path) as db:
        db.execute("DROP TABLE field_options"); db.execute("DROP TABLE table_views")
        db.execute("PRAGMA user_version=13")
    result = TableIO(other).import_file(json_path, format="json", name="Core copy", expected_revision=0)
    assert other.info()["format_version"] == 13 and matrix(other, result["table_id"]) == [[7]]


@pytest.mark.parametrize("value", [1.0000000000000002, 3.141592653589793, 5e-324, -0.0])
def test_excel_preserves_high_precision_floats_subnormals_and_signed_zero_as_text(store, tmp_path, value):
    import math
    ids = [uid() for _ in range(3)]
    store.apply([{"op": "create_table", "id": ids[0], "name": "Precision"},
                 {"op": "add_field", "table_id": ids[0], "id": ids[1], "name": "X", "type": "number"},
                 {"op": "add_record", "table_id": ids[0], "id": ids[2]},
                 {"op": "set_cell", "table_id": ids[0], "record_id": ids[2], "field_id": ids[1], "value": value}], expected_revision=0)
    path = tmp_path / "precise.xlsx"
    service = TableIO(store)
    service.export_file(ids[0], path, format="xlsx", expected_revision=1)
    workbook = openpyxl.load_workbook(path)
    assert workbook.active["A2"].data_type == "s" and workbook.active["A2"].value == json.dumps(value)
    workbook.close()
    copied = service.import_file(path, format="xlsx", name="Exact", expected_revision=1, types={"X": "number"})
    actual = matrix(store, copied["table_id"])[0][0]
    assert actual == value and math.copysign(1, actual) == math.copysign(1, value)
