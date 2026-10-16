"""Bounded CSV, XLSX and typed JSON value exchange; never a project backup.

JSON preserves types/options and absent versus null cells, but deliberately materializes
successful cell definitions. Import creates new identities and keeps object references
bound to their original project/objects. Excel imports never calculate a formula.
"""
from datetime import date, datetime, time
import hashlib
import io
import json
import math
import os
from pathlib import Path
import re
import stat
import sys
import tempfile
from uuid import uuid4
import zipfile

from .csv_io import (MAX_BYTES, MAX_EXPORT_ROWS, MAX_FIELDS, TableCSV, _constant,
                     _headers, _literal, _path)
from .field_types import check, options
from .store import ProjectError, RevisionConflict, _expected_revision, _id, _text

JSON_FORMAT = "stk.table"
JSON_VERSION = 1
MAX_EXPANDED_BYTES = 32 * 1024 * 1024
MAX_ZIP_MEMBERS = 2048
MAX_COMMAND_BYTES = 1024 * 1024
FORMATS = {"csv", "xlsx", "json"}


def _format(value):
    if type(value) is not str or value not in FORMATS:
        raise ProjectError("Table format must be csv, xlsx or json")
    return value


def _raw(path):
    descriptor = os.open(path, os.O_RDONLY | getattr(os, "O_NONBLOCK", 0))
    with os.fdopen(descriptor, "rb") as stream:
        before = os.fstat(stream.fileno())
        if not stat.S_ISREG(before.st_mode):
            raise ProjectError("Table source must be a regular file")
        data = stream.read(MAX_BYTES + 1)
        after = os.fstat(stream.fileno())
    if len(data) > MAX_BYTES:
        raise ProjectError("Table source exceeds 8 MiB")
    identity = lambda s: (s.st_dev, s.st_ino, s.st_size, s.st_mtime_ns, s.st_ctime_ns)
    if identity(before) != identity(after):
        raise ProjectError("Table source changed while reading; preview it again")
    return data, hashlib.sha256(data).hexdigest()


def _json_bytes(value):
    try:
        return json.dumps(value, ensure_ascii=False, allow_nan=False, separators=(",", ":")).encode("utf-8")
    except (UnicodeError, ValueError, TypeError, RecursionError):
        raise ProjectError("Table data must be finite, bounded UTF-8 JSON") from None


def _duplicate_keys(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise ProjectError("Table JSON contains duplicate object keys")
        result[key] = value
    return result


def _bounded_json(data):
    try:
        text = data.decode("utf-8-sig")
        depth, quoted, escaped = 0, False, False
        for char in text:
            if quoted:
                if escaped:
                    escaped = False
                elif char == "\\":
                    escaped = True
                elif char == '"':
                    quoted = False
            elif char == '"':
                quoted = True
            elif char in "[{":
                depth += 1
                if depth > 64:
                    raise ProjectError("Table JSON exceeds depth 64")
            elif char in "]}":
                depth -= 1
        return json.loads(text, parse_constant=_constant, object_pairs_hook=_duplicate_keys)
    except (UnicodeError, ValueError, TypeError, RecursionError):
        raise ProjectError("Table source must be valid finite UTF-8 JSON") from None


def _dependencies():
    try:
        import openpyxl
        from defusedxml import ElementTree
    except ImportError:
        raise ProjectError("Excel exchange requires openpyxl and defusedxml; install the STK table dependencies") from None
    return openpyxl, ElementTree


def _excel_text(cell):
    if cell.data_type == "f":
        raise ProjectError("Excel formula cells are not imported; replace formulas with explicit values first")
    if cell.data_type == "e":
        raise ProjectError("Excel error cells cannot be imported")
    value = cell.value
    if value is None:
        return ""
    if type(value) is str:
        return value
    if type(value) is bool:
        return "true" if value else "false"
    if type(value) in (int, float):
        if type(value) is float and not math.isfinite(value):
            raise ProjectError("Excel numbers must be finite")
        return json.dumps(value, allow_nan=False)
    if isinstance(value, (date, datetime, time)):
        if isinstance(value, datetime) and value.time() == time():
            return value.date().isoformat()
        return value.isoformat()
    raise ProjectError("Unsupported Excel cell value")


def _xlsx(data, sheet):
    openpyxl, safe_xml = _dependencies()
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            names = [m.filename for m in members]
            if (len(members) > MAX_ZIP_MEMBERS or len(set(names)) != len(names) or
                    sum(m.file_size for m in members) > MAX_EXPANDED_BYTES or
                    any(m.flag_bits & 1 for m in members)):
                raise ProjectError("Excel archive exceeds limits, is encrypted or has duplicate entries")
            for member in members:
                lowered = member.filename.lower()
                if "vbaproject" in lowered or lowered.startswith("xl/externallinks/"):
                    raise ProjectError("Excel macros and external workbook links are not supported")
                if lowered.endswith((".xml", ".rels")):
                    raw = archive.read(member)
                    if len(raw) != member.file_size:
                        raise ProjectError("Excel archive member size changed")
                    root = safe_xml.fromstring(raw, forbid_dtd=True, forbid_entities=True, forbid_external=True)
                    for node in root.iter():
                        if lowered.startswith("xl/worksheets/"):
                            tag = node.tag.rsplit("}", 1)[-1]
                            if tag == "row" and "r" in node.attrib:
                                row_number = node.attrib["r"]
                                if (not row_number.isdecimal() or len(row_number) > 5 or
                                        not 1 <= int(row_number) <= MAX_EXPORT_ROWS + 1):
                                    raise ProjectError("Excel worksheet exceeds 10000 data rows")
                            if tag == "c" and "r" in node.attrib:
                                address = re.fullmatch(r"([A-Z]{1,3})([1-9][0-9]{0,4})", node.attrib["r"])
                                if address is None:
                                    raise ProjectError("Excel cell address exceeds worksheet limits")
                                column = 0
                                for letter in address[1]:
                                    column = column * 26 + ord(letter) - ord("A") + 1
                                if column > MAX_FIELDS or int(address[2]) > MAX_EXPORT_ROWS + 1:
                                    raise ProjectError("Excel worksheet exceeds 10000 rows or 64 columns")
                        if node.attrib.get("TargetMode", "").lower() == "external":
                            raise ProjectError("Excel external links are not supported")
                        if "macroenabled" in node.attrib.get("ContentType", "").lower():
                            raise ProjectError("Excel macros are not supported")
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False,
                                         keep_vba=False, keep_links=False)
        try:
            names = workbook.sheetnames
            if not names or len(names) > 256:
                raise ProjectError("Excel workbook needs 1 to 256 worksheets")
            if sheet is not None and (type(sheet) is not str or sheet not in names):
                raise ProjectError("Selected Excel worksheet does not exist")
            selected = names[0] if sheet is None else sheet
            worksheet = workbook[selected]
            if not hasattr(worksheet, "iter_rows"):
                raise ProjectError("Selected Excel sheet is not a worksheet")
            if ((worksheet.max_row is not None and worksheet.max_row > MAX_EXPORT_ROWS + 1) or
                    (worksheet.max_column is not None and worksheet.max_column > MAX_FIELDS)):
                raise ProjectError("Excel worksheet exceeds 10000 rows or 64 columns")
            rows = worksheet.iter_rows()
            first = next(rows, None)
            headers = [_excel_text(cell) for cell in first] if first else None
            _headers(headers)
            values = []
            for cells in rows:
                if len(values) >= MAX_EXPORT_ROWS:
                    raise ProjectError("Excel worksheet exceeds 10000 data rows")
                if len(cells) > len(headers):
                    raise ProjectError("Excel worksheet has values beyond its named columns")
                # Dimensionless streaming sheets contain sparse/empty rows. Their missing
                # trailing cells are blank, not an invalid CSV-shaped row.
                values.append([_excel_text(cell) for cell in cells] + [""] * (len(headers) - len(cells)))
            return {"columns": [{"name": h, "type": "text", "unit": None} for h in headers],
                    "rows": values, "sheets": [{"name": n} for n in names], "sheet": selected,
                    "schema_present": False, "table_name": None,
                    "warnings": ["Excel imports values only; choose explicit column types and units."]}
        finally:
            workbook.close()
    except ProjectError:
        raise
    except Exception as exc:
        # Parser errors may include workbook contents/paths; return a bounded generic diagnosis.
        raise ProjectError(f"Invalid or unsupported Excel workbook ({type(exc).__name__})") from None


def _json_table(data, project_id):
    document = _bounded_json(data)
    if (type(document) is not dict or set(document) != {
            "format", "version", "mode", "source", "name", "fields", "records"} or
            document["format"] != JSON_FORMAT or type(document["version"]) is not int or
            document["version"] != JSON_VERSION or document["mode"] != "evaluated_values"):
        raise ProjectError("JSON requires the stk.table version 1 evaluated_values format")
    source = document["source"]
    if type(source) is not dict or set(source) != {"project_id", "table_id", "revision"}:
        raise ProjectError("Table JSON source descriptor is invalid")
    _id(source["project_id"]); _id(source["table_id"]); _expected_revision(source["revision"])
    _text(document["name"], "Table JSON name")
    fields = document["fields"]
    if type(fields) is not list or not 1 <= len(fields) <= MAX_FIELDS:
        raise ProjectError("Table JSON requires 1 to 64 fields")
    identities = set()
    for field in fields:
        if type(field) is not dict or set(field) != {"id", "name", "type", "unit", "options"}:
            raise ProjectError("Invalid Table JSON field descriptor")
        _id(field["id"])
        if field["id"] in identities:
            raise ProjectError("Table JSON field IDs must be unique")
        identities.add(field["id"])
        field["options"] = options(field["type"], field["options"])
        if field["unit"] is not None:
            _text(field["unit"], "Unit")
            if field["type"] not in {"integer", "number"}:
                raise ProjectError("Only numeric fields may declare units")
    _headers([f["name"] for f in fields])
    records = document["records"]
    if type(records) is not list or len(records) > MAX_EXPORT_ROWS:
        raise ProjectError("Table JSON supports at most 10000 records")
    record_ids, rows, commands = set(), [], 1 + len(fields)
    for record in records:
        if type(record) is not dict or set(record) != {"id", "values"}:
            raise ProjectError("Table JSON record requires id and values")
        _id(record["id"])
        if record["id"] in record_ids:
            raise ProjectError("Table JSON record IDs must be unique")
        record_ids.add(record["id"])
        values = record["values"]
        if type(values) is not dict or values.keys() - identities:
            raise ProjectError("Table JSON values contain unknown fields")
        for field in fields:
            if field["id"] in values:
                value = values[field["id"]]
                check(value, field["type"], field["options"], project_id)
                if field["type"] == "object" and value is not None and source["project_id"] != project_id:
                    raise ProjectError("Object references can only be imported into their original project")
        commands += 1 + len(values)
        rows.append([values.get(f["id"]) for f in fields])
    return {"columns": [{k: f[k] for k in ("name", "type", "unit", "options")} for f in fields],
            "rows": rows, "sheets": [], "sheet": None, "schema_present": True,
            "table_name": document["name"], "document": document, "command_count": commands,
            "warnings": ["JSON restores typed values; cell formulas, views and original identities are not restored."]}


def _publish(store, destination, data):
    if len(data) > MAX_BYTES:
        raise ProjectError("Table export exceeds 8 MiB")
    destination, temporary = _path(store, destination), None
    try:
        with tempfile.NamedTemporaryFile(dir=destination.parent, prefix=".stk-table-", delete=False) as stream:
            temporary = Path(stream.name)
            stream.write(data)
            stream.flush()
            os.fsync(stream.fileno())
        os.link(temporary, destination)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    return {"path": str(destination), "size": len(data), "sha256": hashlib.sha256(data).hexdigest()}


class TableIO:
    def __init__(self, store):
        self.store = store

    def _revision(self, expected_revision):
        _expected_revision(expected_revision)
        current = self.store.info()
        if current["revision"] != expected_revision:
            raise RevisionConflict("Table exchange revision is stale")
        return current

    def _load(self, source, format, sheet):
        data, digest = _raw(_path(self.store, source))
        if format == "xlsx":
            loaded = _xlsx(data, sheet)
        else:
            if sheet is not None:
                raise ProjectError("Only Excel supports a worksheet selection")
            loaded = _json_table(data, self.store.info()["id"])
        return loaded, digest

    def preview(self, source, *, format, expected_revision, delimiter=",", sheet=None, limit=20):
        format = _format(format)
        if type(limit) is not int or not 1 <= limit <= 100:
            raise ProjectError("Table preview limit must be 1 to 100")
        self._revision(expected_revision)
        if format == "csv":
            if sheet is not None:
                raise ProjectError("Only Excel supports a worksheet selection")
            result = TableCSV(self.store).preview(source, expected_revision=expected_revision,
                                                 delimiter=delimiter, limit=limit)
            return {**result, "format": format, "sheets": [], "sheet": None,
                    "schema_present": False, "table_name": None, "warnings": []}
        loaded, digest = self._load(source, format, sheet)
        self._revision(expected_revision)
        count = len(loaded["rows"])
        commands = loaded.get("command_count", 1 + len(loaded["columns"]) + count * (1 + len(loaded["columns"])))
        can_import = commands <= 1000
        if can_import:
            plan, _, _, _ = self._commands(loaded, format, loaded["table_name"] or "Imported table",
                                           None, None, self.store.info()["id"])
            # Reserve the maximum UTF-8 table-name length, because the user can rename
            # the destination after this read-only preview.
            can_import = len(_json_bytes(plan)) + 1024 <= MAX_COMMAND_BYTES
        return {k: v for k, v in loaded.items() if k not in {"document", "rows"}} | {
            "revision": expected_revision, "format": format, "source_sha256": digest,
            "rows": loaded["rows"][:limit], "total_rows": count, "truncated": count > limit,
            "command_count": commands, "can_import": can_import}

    def import_file(self, source, *, format, name, expected_revision, types=None, units=None,
                    delimiter=",", sheet=None, expected_sha256=None):
        format = _format(format)
        self._revision(expected_revision)
        _text(name, "Table name")
        if format == "csv":
            if sheet is not None:
                raise ProjectError("Only Excel supports a worksheet selection")
            return TableCSV(self.store).import_file(source, name=name, expected_revision=expected_revision,
                types=types, units=units, delimiter=delimiter, expected_sha256=expected_sha256) | {"format": format}
        loaded, digest = self._load(source, format, sheet)
        project_id = self.store.info()["id"]
        if expected_sha256 is not None and expected_sha256 != digest:
            raise RevisionConflict("Table source changed since its preview; preview it again")
        commands, table_id, field_ids, record_ids = self._commands(
            loaded, format, name, types, units, project_id)
        if len(_json_bytes(commands)) > MAX_COMMAND_BYTES:
            raise ProjectError("Table import exceeds the 1 MiB command limit")
        result = self.store.apply(commands, expected_revision=expected_revision)
        return {"revision": result["revision"], "format": format, "table_id": table_id,
                "field_ids": field_ids, "record_ids": record_ids, "rows": len(record_ids),
                "columns": len(field_ids), "source_sha256": digest, "sheet": loaded["sheet"]}

    @staticmethod
    def _commands(loaded, format, name, types, units, project_id):
        types, units = {} if types is None else types, {} if units is None else units
        if type(types) is not dict or type(units) is not dict:
            raise ProjectError("Table types and units must be column-name mappings")
        if format == "json" and (types or units):
            raise ProjectError("Typed JSON contains its own schema; type/unit overrides are not supported")
        headers = [f["name"] for f in loaded["columns"]]
        if set(types) - set(headers) or set(units) - set(headers):
            raise ProjectError("Types/units contain columns absent from the table")
        table_id, field_ids = str(uuid4()), [str(uuid4()) for _ in headers]
        commands = [{"op": "create_table", "id": table_id, "name": name}]
        fields = []
        for column, field_id in zip(loaded["columns"], field_ids):
            kind, unit = types.get(column["name"], column["type"]), units.get(column["name"], column["unit"])
            config = options(kind, column.get("options"))
            if unit is not None:
                _text(unit, "Unit")
                if kind not in {"integer", "number"}:
                    raise ProjectError("Only numeric columns may declare units")
            command = {"op": "add_field", "table_id": table_id, "id": field_id,
                       "name": column["name"], "type": kind, "unit": unit}
            # Core fields remain importable into format 13; format 14 owns new types/options.
            if kind not in {"text", "integer", "number", "boolean", "json"}:
                command["options"] = config
            commands.append(command)
            fields.append({"id": field_id, "type": kind, "options": config})
        record_ids = []
        for index, row in enumerate(loaded["rows"]):
            identity = str(uuid4())
            record_ids.append(identity)
            commands.append({"op": "add_record", "table_id": table_id, "id": identity})
            for column_index, (field_id, field, value) in enumerate(zip(field_ids, fields, row)):
                if format == "json":
                    original = loaded["document"]["records"][index]["values"]
                    if loaded["document"]["fields"][column_index]["id"] not in original:
                        continue
                else:
                    value = _literal(value, field["type"], index + 1, headers[column_index])
                check(value, field["type"], field["options"], project_id)
                commands.append({"op": "set_cell", "table_id": table_id, "record_id": identity,
                                 "field_id": field_id, "value": value})
            if len(commands) > 1000:
                raise ProjectError("Table import exceeds one atomic edit batch (1000 commands)")
        return commands, table_id, field_ids, record_ids

    def _table(self, table_id, expected_revision, definition, view_id):
        snapshot = self.store.snapshot()
        if snapshot["project"]["revision"] != expected_revision:
            raise RevisionConflict("Table export revision is stale")
        table = next((t for t in snapshot["tables"] if t["id"] == table_id), None)
        if table is None:
            raise ProjectError("Table does not exist")
        if definition is not None or view_id is not None:
            records, page = [], self.store.table.query(table_id, expected_revision=expected_revision,
                definition=definition, view_id=view_id, limit=500)
            if page["issues"] or page["total"] > MAX_EXPORT_ROWS:
                raise ProjectError("Cannot export an unresolved view or more than 10000 records")
            hidden = {c["field_id"] for c in page["definition"]["columns"] if c["hidden"]}
            fields = [f for f in page["fields"] if f["id"] not in hidden]
            while True:
                records.extend(page["records"])
                if page["next_offset"] is None:
                    break
                page = self.store.table.query(table_id, expected_revision=expected_revision,
                    definition=definition, view_id=view_id, offset=page["next_offset"], limit=500)
            table = {**table, "fields": fields, "records": records}
        _headers([f["name"] for f in table["fields"]])
        if len(table["records"]) > MAX_EXPORT_ROWS:
            raise ProjectError("Table export supports at most 10000 records")
        for record in table["records"]:
            for field in table["fields"]:
                if record.get("evaluations", {}).get(field["id"], {}).get("state") == "error":
                    raise ProjectError(f"Cannot export formula error at record {record['id']}, field {field['id']}")
        return snapshot["project"], table

    def export_file(self, table_id, destination, *, format, expected_revision, definition=None,
                    view_id=None, delimiter=","):
        format = _format(format)
        self._revision(expected_revision)
        if format == "csv":
            return TableCSV(self.store).export_file(table_id, destination, expected_revision=expected_revision,
                definition=definition, view_id=view_id, delimiter=delimiter) | {"format": format}
        project, table = self._table(table_id, expected_revision, definition, view_id)
        fields, records = table["fields"], table["records"]
        warnings = ["Cell formulas are exported as evaluated values."]
        if format == "json":
            field_ids = {f["id"] for f in fields}
            document = {"format": JSON_FORMAT, "version": JSON_VERSION, "mode": "evaluated_values",
                "source": {"project_id": project["id"], "table_id": table_id, "revision": expected_revision},
                "name": table["name"],
                "fields": [{**f, "options": f.get("options", {"version": 1})} for f in fields],
                "records": [{"id": r["id"], "values": {k: v for k, v in r["values"].items() if k in field_ids}} for r in records]}
            data = _json_bytes(document)
        else:
            openpyxl, _ = _dependencies()
            from openpyxl.cell import WriteOnlyCell
            all_rows = [[f["name"] for f in fields]]
            for record in records:
                values = []
                for field in fields:
                    value = record["values"].get(field["id"])
                    if value is not None and field["type"] in {"json", "multi_select", "object"}:
                        value = _json_bytes(value).decode("utf-8")
                    elif type(value) is int and len(str(abs(value))) > 15:
                        value = str(value)  # Excel's 15-digit precision cannot represent int64.
                    elif type(value) is float and (float(f"{value:.15g}") != value or
                            0 < abs(value) < sys.float_info.min or value == 0 and math.copysign(1, value) < 0):
                        # Preserve the actual scientific double, including subnormal values
                        # and signed zero, when Excel's numeric cell cannot round-trip it.
                        value = json.dumps(value, allow_nan=False)
                    values.append(value)
                all_rows.append(values)
            text_bytes = 0
            for row in all_rows:
                for value in row:
                    if type(value) is str:
                        try:
                            text_bytes += len(value.encode("utf-8"))
                        except UnicodeError:
                            raise ProjectError("Excel text must be valid UTF-8") from None
                        if len(value) > 32767 or re.search(r"[\x00-\x08\x0b-\x0c\x0e-\x1f]", value):
                            raise ProjectError("Excel strings must fit 32767 characters and contain no XML controls")
            if text_bytes > MAX_EXPANDED_BYTES:
                raise ProjectError("Excel export text exceeds 32 MiB")
            workbook = openpyxl.Workbook(write_only=True)
            worksheet = workbook.create_sheet("Table")
            worksheet.freeze_panes = "A2"
            try:
                for values in all_rows:
                    cells = []
                    for value in values:
                        cell = WriteOnlyCell(worksheet, value=value)
                        if type(value) is str:
                            cell.data_type = "s"  # =, +, -, @ strings are inert text.
                        cells.append(cell)
                    worksheet.append(cells)
                output = io.BytesIO()
                workbook.save(output)
                data = output.getvalue()
                with zipfile.ZipFile(io.BytesIO(data)) as archive:
                    if sum(m.file_size for m in archive.infolist()) > MAX_EXPANDED_BYTES:
                        raise ProjectError("Excel export exceeds its 32 MiB expanded-file limit")
            finally:
                workbook.close()
            warnings.append("Excel omits STK field schemas; map column types/units on import. High precision numbers, subnormal floats and signed zero are text.")
        self._revision(expected_revision)
        published = _publish(self.store, destination, data)
        return {"revision": expected_revision, "table_id": table_id, "format": format,
                "rows": len(records), "columns": len(fields), "warnings": warnings, **published}
